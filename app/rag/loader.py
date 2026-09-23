"""
Unified document loader for PDF / DOCX / DOC / TXT / Markdown.

R130: every loader below ends in :func:`sanitize_text`, because pypdf really does hand
back a NUL character (measured in documents/AI-Agent学习路线图.pdf) and a PostgreSQL
text column cannot hold one.
"""
from pathlib import Path

from pypdf import PdfReader

from app.common.logger import logger

#: U+0000 -- the one character a PostgreSQL ``text`` column refuses, and therefore the
#: one character this module drops. psycopg only complains about it from inside
#: ``executemany``, by which point the batch has no document name left to report, so
#: the drop belongs here, at the edge where a third-party parser stops being our problem.
#: The mirror in app/rag/pg_store.py keeps its own gate for anything arriving by another
#: route: two layers, each doing its own job.
NUL_CHARACTER = "\x00"


def sanitize_text(text: str) -> str:
    """Drop NUL characters and nothing else.

    Deliberately surgical (R130 判据①): whitespace, blank lines, emoji, letter case and
    every other control character belong to the document, so they leave unchanged. A text
    without NUL comes back byte-identical, which is what keeps the already-indexed corpus
    (R130 判据④: 985 chunks, none of them dirty) untouched by this ticket.
    """
    if NUL_CHARACTER not in text:
        return text
    return text.replace(NUL_CHARACTER, "")


def load_pdf(file_path: str) -> str:
    """Extract PDF text locally with pypdf."""
    reader = PdfReader(file_path)
    parts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(page_text.strip())
    text = sanitize_text("\n\n".join(parts).strip())
    logger.info(f"Loaded PDF with pypdf: {file_path} ({len(text)} chars)")
    return text


def load_docx(file_path: str) -> str:
    """Extract plain text from .docx."""
    from docx import Document

    doc = Document(file_path)
    return sanitize_text("\n".join(p.text for p in doc.paragraphs if p.text.strip()))


def load_doc(file_path: str) -> str:
    """Best-effort extraction for old .doc files."""
    import olefile

    ole = olefile.OleFileIO(file_path)
    if ole.exists("WordDocument"):
        stream = ole.openstream("WordDocument")
        raw = stream.read()
        text = "".join(chr(b) for b in raw if 31 < b < 127 or b in (10, 13))
        lines = [line.strip() for line in text.split("\n") if len(line.strip()) > 2]
        ole.close()
        return sanitize_text("\n".join(lines))
    ole.close()
    return ""


def load_txt(file_path: str) -> str:
    """Read TXT with automatic encoding detection."""
    for encoding in ["utf-8", "gbk", "gb2312"]:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return sanitize_text(f.read())
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Unable to detect text encoding: {file_path}")


def load_md(file_path: str) -> str:
    """Read Markdown as plain source text, reusing TXT encoding detection."""
    return sanitize_text(load_txt(file_path))


def load_document(file_path: str) -> str:
    """Auto-detect file type and return plain text."""
    ext = Path(file_path).suffix.lower()
    logger.info(f"Loading document: {file_path} ({ext})")

    if ext == ".pdf":
        return sanitize_text(load_pdf(file_path))
    if ext == ".docx":
        return sanitize_text(load_docx(file_path))
    if ext == ".doc":
        return sanitize_text(load_doc(file_path))
    if ext == ".txt":
        return sanitize_text(load_txt(file_path))
    if ext == ".md":
        return sanitize_text(load_md(file_path))
    raise ValueError(f"Unsupported file format: {ext}")

# ---------------------------------------------------------------------------
# V2 structured parsing entry point.
#
# The ``load_*`` functions above stay the plain-text path and are unchanged so
# every existing caller (upload API, preview, tests) keeps its exact behaviour.
# ``parse_document`` below is the richer entry that reports through
# ``DocumentParseResult`` and assembles text, OCR, tables and spreadsheet
# import blocks into one shape without flattening tables into prose.
# ---------------------------------------------------------------------------
from app.rag.parser_result import (
    BlockType,
    ContentBlock,
    DocumentParseResult,
    ParseError,
    ParseStatus,
    ParserType,
    SourceLocation,
)


def _block_has_content(block: ContentBlock) -> bool:
    if block.block_type in (BlockType.TEXT, BlockType.OCR):
        return bool(block.content.strip())
    return bool(block.rows) or bool(block.header)


def _classify_parse(blocks, errors, warnings):
    """Derive status/failure from blocks and top-level errors.

    A document with content and no errors is success; content plus errors is
    partial success; errors with no content is failed; no content and no errors
    is a truthful empty document.
    """
    content = any(_block_has_content(block) for block in blocks)
    if errors and not content:
        return ParseStatus.FAILED, errors[0].message, tuple(warnings), tuple(errors)
    if errors:
        return ParseStatus.PARTIAL_SUCCESS, None, tuple(warnings), tuple(errors)
    if content:
        return ParseStatus.SUCCESS, None, tuple(warnings), tuple(errors)
    return ParseStatus.EMPTY, None, tuple(warnings), tuple(errors)


def _parsers(blocks) -> tuple[ParserType, ...]:
    return tuple(dict.fromkeys(block.parser_type for block in blocks))


def _wrap_plain(file_path: str, ext: str, parser_type: ParserType, text: str) -> DocumentParseResult:
    """Wrap an existing plain-text load into the unified result."""
    source = str(Path(file_path))
    if text:
        block = ContentBlock(
            block_type=BlockType.TEXT,
            parser_type=parser_type,
            content=sanitize_text(text),
            location=SourceLocation(file_path=source),
        )
        return DocumentParseResult.success(
            source, ext, blocks=[block], parsers=[parser_type]
        )
    return DocumentParseResult.empty(source, ext, parsers=[parser_type])


def parse_pdf_result(
    file_path: str,
    *,
    include_tables: bool = True,
    include_ocr: bool = False,
    ocr_engine=None,
    render_page=None,
    confidence_threshold: float | None = None,
) -> DocumentParseResult:
    """Parse a PDF into text / table / OCR blocks with page provenance."""
    source = str(Path(file_path))
    try:
        reader = PdfReader(file_path)
    except Exception as exc:
        error = ParseError(
            code="pdf_open_failed",
            message=f"pypdf could not open PDF: {exc}",
            parser=ParserType.PYPDF,
            location=SourceLocation(file_path=source),
        )
        return DocumentParseResult.failed(
            source,
            ".pdf",
            error.message,
            parsers=[ParserType.PYPDF],
            errors=[error],
        )

    blocks: list[ContentBlock] = []
    errors: list[ParseError] = []
    warnings: list[str] = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:
            errors.append(
                ParseError(
                    code="pdf_page_extract_failed",
                    message=f"page {page_number}: {exc}",
                    parser=ParserType.PYPDF,
                    location=SourceLocation(file_path=source, page=page_number),
                )
            )
            continue
        if page_text.strip():
            blocks.append(
                ContentBlock(
                    block_type=BlockType.TEXT,
                    parser_type=ParserType.PYPDF,
                    content=sanitize_text(page_text.strip()),
                    location=SourceLocation(file_path=source, page=page_number),
                )
            )

    if include_tables:
        from app.rag.table_parser import extract_pdf_tables

        run = extract_pdf_tables(file_path)
        blocks.extend(run.blocks)
        errors.extend(run.errors)
        warnings.extend(run.warnings)

    if include_ocr:
        if render_page is None:
            errors.append(
                ParseError(
                    code="ocr_renderer_missing",
                    message="OCR requested but no page renderer was provided",
                    parser=ParserType.OCR,
                    location=SourceLocation(file_path=source),
                )
            )
        if ocr_engine is None:
            errors.append(
                ParseError(
                    code="ocr_engine_missing",
                    message="OCR requested but no engine was provided",
                    parser=ParserType.OCR,
                    location=SourceLocation(file_path=source),
                )
            )
        if render_page is not None and ocr_engine is not None:
            from app.rag.ocr.pdf_ocr import ocr_scanned_pages

            run = ocr_scanned_pages(
                reader,
                ocr_engine,
                render_page,
                source=source,
                confidence_threshold=confidence_threshold,
            )
            blocks.extend(run.blocks)
            warnings.extend(run.warnings)

    status, failure_reason, warnings, errors = _classify_parse(blocks, errors, warnings)
    return DocumentParseResult(
        source=source,
        extension=".pdf",
        status=status,
        blocks=tuple(blocks),
        parsers=_parsers(blocks),
        failure_reason=failure_reason,
        warnings=warnings,
        errors=errors,
    )


def parse_docx_result(file_path: str, *, include_tables: bool = True) -> DocumentParseResult:
    """Parse a DOCX into paragraph text and structured table blocks."""
    source = str(Path(file_path))
    try:
        from docx import Document

        document = Document(file_path)
        text = sanitize_text("\n".join(p.text for p in document.paragraphs if p.text.strip()))
    except Exception as exc:
        error = ParseError(
            code="docx_open_failed",
            message=f"python-docx could not open DOCX: {exc}",
            parser=ParserType.PYTHON_DOCX,
            location=SourceLocation(file_path=source),
        )
        return DocumentParseResult.failed(
            source,
            ".docx",
            error.message,
            parsers=[ParserType.PYTHON_DOCX],
            errors=[error],
        )

    blocks: list[ContentBlock] = []
    errors: list[ParseError] = []
    warnings: list[str] = []
    if text:
        blocks.append(
            ContentBlock(
                block_type=BlockType.TEXT,
                parser_type=ParserType.PYTHON_DOCX,
                content=text,
                location=SourceLocation(file_path=source),
            )
        )

    if include_tables:
        from app.rag.table_parser import extract_docx_tables

        run = extract_docx_tables(file_path)
        blocks.extend(run.blocks)
        errors.extend(run.errors)
        warnings.extend(run.warnings)

    status, failure_reason, warnings, errors = _classify_parse(blocks, errors, warnings)
    return DocumentParseResult(
        source=source,
        extension=".docx",
        status=status,
        blocks=tuple(blocks),
        parsers=_parsers(blocks),
        failure_reason=failure_reason,
        warnings=warnings,
        errors=errors,
    )


def _spreadsheet_result(source: str, ext: str, run) -> DocumentParseResult:
    failure_reason = run.errors[0].message if run.status is ParseStatus.FAILED and run.errors else None
    return DocumentParseResult(
        source=source,
        extension=ext,
        status=run.status,
        blocks=run.blocks,
        parsers=_parsers(run.blocks),
        failure_reason=failure_reason,
        warnings=run.warnings,
        errors=run.errors,
        metadata=run.metadata,
    )


def parse_document(
    file_path: str,
    *,
    include_tables: bool = True,
    include_ocr: bool = False,
    ocr_engine=None,
    render_page=None,
    confidence_threshold: float | None = None,
    max_rows_per_sheet: int | None = 1000,
) -> DocumentParseResult:
    """Unified structured parse entry point.

    This supersedes ``load_document`` for new callers that want provenance and
    structured blocks.  ``load_document`` is kept intact as the plain-text
    compatibility path.
    """
    ext = Path(file_path).suffix.lower()
    logger.info(f"Parsing document: {file_path} ({ext})")

    if ext == ".pdf":
        return parse_pdf_result(
            file_path,
            include_tables=include_tables,
            include_ocr=include_ocr,
            ocr_engine=ocr_engine,
            render_page=render_page,
            confidence_threshold=confidence_threshold,
        )
    if ext == ".docx":
        return parse_docx_result(file_path, include_tables=include_tables)
    if ext == ".doc":
        return _wrap_plain(file_path, ext, ParserType.OLEFILE, load_doc(file_path))
    if ext == ".txt":
        return _wrap_plain(file_path, ext, ParserType.TEXT_ENCODING, load_txt(file_path))
    if ext == ".md":
        return _wrap_plain(file_path, ext, ParserType.MARKDOWN, load_md(file_path))
    if ext in (".xlsx", ".xls"):
        from app.rag.excel_kb_parser import extract_excel_kb

        run = extract_excel_kb(file_path, max_rows_per_sheet=max_rows_per_sheet)
        return _spreadsheet_result(str(Path(file_path)), ext, run)
    if ext == ".csv":
        from app.rag.excel_kb_parser import extract_csv_kb

        run = extract_csv_kb(file_path, max_rows=max_rows_per_sheet)
        return _spreadsheet_result(str(Path(file_path)), ext, run)
    raise ValueError(f"Unsupported file format: {ext}")
