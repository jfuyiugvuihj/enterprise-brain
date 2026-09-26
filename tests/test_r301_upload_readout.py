r"""R301 判据①②③④⑤⑥ · 上传当场把「这份 PDF 是怎么读出来的」交回调用方。

病灶：``load_document`` 对外只交正文串，而 loader 内部早就组装好的
``DocumentExtraction``（``pdf`` 逐页来源账 + ``tables`` 表格账）除了测试没人消费 ——
客户传一份扫描件，屏幕上永远看不到「一共几页 / 几页是扫的 / 哪几页 OCR 没跑成」。

本件钉的格（全程离线：栅格化用桩、引擎用 R298 那套假对象、不连库、不起服务）：
① 响应新增一格 ``pdf_extraction``，承载 page_count / 扫描页数 / ocr_attempted /
   ocr_available / 逐页来源账 / 没跑成的页号，且只报读数不引申；
② 降级句只抄 ``report.degradation_sentence`` 那一把尺，「引擎不可用」与「引擎在、这一页
   没跑成」在响应里分得开；③ 非 PDF（.txt/.md/.docx）整格 ``None``，不塞空对象；
④ 这一格不落库（不写目录行、不加迁移）；⑤ 绝对路径不外泄；⑥ 端到端凭据：真走
   ``upload_document``，「扫描页 3/6」真的出现在响应 JSON 里；⑦ 反证五把，见文件末尾。

为什么上传路径只多调一枚入口、不「先 load_document 再补一遍账」：那等于把 R298 的 OCR 在
客户机上重跑一次。``test_one_pdf_is_parsed_exactly_once`` 钉的就是这条。
"""
from __future__ import annotations

import asyncio
import contextlib
import io
import json
import logging
import re
from pathlib import Path

import numpy as np
import pytest

from app.rag import loader
from app.rag import ocr as ocr_channel
from app.rag.loader import PAGE_SOURCE_BLANK, PAGE_SOURCE_OCR, PAGE_SOURCE_OCR_DEGRADED
from app.rag.loader import PAGE_SOURCE_OCR_EMPTY, PAGE_SOURCE_TEXT_LAYER

REPO_ROOT = Path(__file__).resolve().parents[1]
MIXED_PDF = REPO_ROOT / "tests" / "fixtures" / "r298_mixed_pages.pdf"
BLANK_PDF = REPO_ROOT / "tests" / "fixtures" / "r298_blank.pdf"

#: 混合件逐页真值（现取 2026-09-26，与 tests/test_r298_scan_detection.py 同一份）：
#: 6 页里 2/4/5 三页是扫描页，第 3 页「文本层 + 一张插图」不该 OCR，第 6 页真空白。
MIXED_PAGE_COUNT = 6
MIXED_SCANNED_PAGES = [2, 4, 5]
CELL_KEY = "pdf_extraction"


@contextlib.contextmanager
def _quiet_parser_warnings():
    """pypdf 解析真语料会往 stderr 刷坏 CMap 行（R130/R298 同一手法），本件不修它，只静音。"""
    pypdf_logger = logging.getLogger("pypdf")
    previous = pypdf_logger.level
    pypdf_logger.setLevel(logging.ERROR)
    try:
        with contextlib.redirect_stderr(io.StringIO()):
            yield
    finally:
        pypdf_logger.setLevel(previous)


class PageAwareEngine:
    """按页给答案的假引擎（R298 的手法）：页码从桩栅格的形状里解回来。

    ``no_text_pages`` 里的页「跑成了、图里确实没字」（``ocr-empty``），``fail_pages`` 里的页
    「这一页没跑成」（``ocr-degraded``）—— 两态在 loader 里本来就是两句话，本件要它们在响应里
    也还是两句话。
    """

    def __init__(self, no_text_pages=(), fail_pages=()) -> None:
        self.no_text_pages = set(no_text_pages)
        self.fail_pages = set(fail_pages)
        self.calls = 0
        self.seen_pages: list[int] = []

    def __call__(self, array):  # noqa: ANN001 - 与 RapidOCR.__call__ 同形
        page_number = int(array.shape[0]) - 8
        self.seen_pages.append(page_number)
        self.calls += 1
        if page_number in self.fail_pages:
            raise RuntimeError("假引擎在这一页炸了")
        if page_number in self.no_text_pages:
            return (None, [0.0])
        box = [[0.0, 0.0], [100.0, 0.0], [100.0, 20.0], [0.0, 20.0]]
        return ([[box, f"扫描页正文 第{page_number}页", 0.99]], [0.0])


@pytest.fixture(autouse=True)
def stub_raster(monkeypatch):
    """栅格化桩：回一张高里编着页码的小图，于是假引擎能按页给不同答案，不读真像素。"""

    def _fake(pdf_path: str, page_number: int, *, dpi: int):
        return np.zeros((page_number + 8, page_number + 16, 3), dtype="uint8")

    monkeypatch.setattr(ocr_channel, "rasterize_page", _fake)


def _engine_available(monkeypatch, engine: PageAwareEngine) -> PageAwareEngine:
    """把假引擎接到 ``ocr_pages`` 的取引擎那一口上（生产调用不传 engine，走这里）。"""
    monkeypatch.setattr(ocr_channel, "unavailable_reason", lambda: None)
    monkeypatch.setattr(ocr_channel, "get_engine", lambda **kwargs: engine)
    return engine


class RecordingRetriever:
    """只回答预置结果并记录被调用过没有（与 tests/test_r49_upload_contract.py 同一套桩）。"""

    def __init__(self, result=(True, "已添加 2 个文本块")):
        self.result = result
        self.add_calls = []

    def add_document(self, filename, content, classification, department):
        self.add_calls.append(filename)
        return self.result

    def list_documents(self):
        return list(self.add_calls)


def _wire(monkeypatch, tmp_path, retriever=None):
    """起一具离线的上传装置。

    🔴 与 R49 那具桩唯一的关键差别：这里**不**替换 ``chat.load_document``。本件的凭据必须
    走真的 loader，否则「扫描页 3/6」是桩话 —— 被换掉的 loader 产不出账，那一格只会是 None。
    """
    from app.api.v1 import chat
    from app.documents import catalog

    monkeypatch.setattr(chat, "DOCUMENTS_DIR", str(tmp_path))
    monkeypatch.setattr(catalog, "DOCUMENTS_DIR", str(tmp_path), raising=False)
    monkeypatch.setattr(catalog, "_database_available", lambda: False, raising=False)
    monkeypatch.setattr(chat, "catalog_database_available", lambda: False)
    monkeypatch.setattr(chat, "retriever", retriever or RecordingRetriever())
    monkeypatch.setattr(chat, "peek_next_document_version", lambda filename: 1)
    return chat


def _upload(chat, path: Path, filename: str | None = None):
    from fastapi import UploadFile

    body = path.read_bytes()
    with _quiet_parser_warnings():
        return asyncio.run(
            chat.upload_document(
                UploadFile(filename=filename or path.name, file=io.BytesIO(body)),
                classification=1,
                department="",
            )
        )


def _mini_docx(directory: Path) -> Path:
    """一枚真能开的 .docx（R304 同一手法）：判据③ 要的是「非 PDF 那一格怎么呈现」。"""
    from docx import Document

    directory.mkdir(parents=True, exist_ok=True)
    target = directory / "notes.docx"
    document = Document()
    document.add_paragraph("费用报销管理制度正文。员工发生的差旅费须在三十日内提交审批。" * 4)
    document.save(str(target))
    return target


# ------------------------------------------------------------------ 三把共用的尺
# 判据①③⑤ 各有一条"不许"。它们写成可复用的校验函数，绿件用一次、反证件再故意违反一次：
# 于是"摘了这一把 ⇒ 红的正是这一条"是当场可证的，不是一句承诺。

#: 绝对路径的形状（Windows 盘符 + 分隔符，或 POSIX 绝对路径开头两段）。
_ABS_PATH = re.compile(r"[A-Za-z]:[\\/]|^/[A-Za-z0-9_.-]+/")
_FORBIDDEN_KEYS = {"file_path", "storage_path", "abs_path", "local_path", "path"}
_REQUIRED_READINGS = (
    "page_count",
    "scanned_pages",
    "ocr_attempted",
    "ocr_available",
    "source_counts",
    "ocr_degraded_page_numbers",
)
# 判据①：「这一页 OCR 没跑成」与「这一页没有内容」是两件事，响应里不许出现后者那种引申。
_EMPTINESS_CLAIMS = ("没有内容", "无内容", "没有文字", "没有字", "空白页", "内容为空")
#: 词表只有 loader 那一套（判据①：别新造第二套）。
_SOURCE_VOCABULARY = {
    PAGE_SOURCE_TEXT_LAYER,
    PAGE_SOURCE_OCR,
    PAGE_SOURCE_OCR_EMPTY,
    PAGE_SOURCE_OCR_DEGRADED,
    PAGE_SOURCE_BLANK,
}
#: 扫描页 = 这三枚来源码的页（有文本层的页与空白页都不在 OCR 通道里）。
_SCANNED_SOURCES = (PAGE_SOURCE_OCR, PAGE_SOURCE_OCR_EMPTY, PAGE_SOURCE_OCR_DEGRADED)
#: 那一格的键集，与 docs/api/contract-v1.md 的 R301 那张表逐条对齐。
#: 四条返回路径（已入索引 / 被策略排除 / 重复未收 / 索引拒收但保留文件）都必须交出同一集键。
_CELL_KEYS = {
    "page_count",
    "scanned_pages",
    "scanned_page_numbers",
    "ocr_attempted",
    "ocr_available",
    "ocr_engine",
    "ocr_dpi",
    "source_counts",
    "ocr_degraded_page_numbers",
    "degradation_note",
}


def _assert_reading_or_nothing(cell):
    """判据③ 的形状尺：要么带着真读数，要么整格 None —— 没有第三种。

    空 dict / page_count=0 的账都算"冒充读过了"，一律判红：非 PDF 那一支今天什么都没读，
    它就该长成"没有这一格读数"的样子。
    """
    if cell is None:
        return
    assert isinstance(cell, dict), f"这一格要么是读数对象要么是 None：{type(cell)}"
    assert cell, "非 PDF 塞了一个空对象：那读起来像「查过了，什么都没查到」"
    for key in _REQUIRED_READINGS:
        assert key in cell, f"这一格少了读数 {key}"
    assert cell["page_count"] > 0, "page_count=0 的账不是读数，是猜的"
    assert sum(cell["source_counts"].values()) == cell["page_count"], "逐页来源账对不上总页数"
    assert set(cell["source_counts"]) <= _SOURCE_VOCABULARY, "来源码不在 loader 的词表里"
    from_sources = sum(cell["source_counts"].get(source, 0) for source in _SCANNED_SOURCES)
    assert from_sources == cell["scanned_pages"], "扫描页读数与逐页来源账自相矛盾"


def _walk(node):
    """把响应里的每一枚键与每一个值都交出来：判据⑤ 要扫的是整格，只看顶层会被嵌套绕过。"""
    if isinstance(node, dict):
        for key, value in node.items():
            yield "key", key
            yield from _walk(value)
    elif isinstance(node, (list, tuple)):
        for item in node:
            yield from _walk(item)
    else:
        yield "value", node


def _assert_no_server_path(response, tmp_path: Path):
    """判据⑤：客户只许看见自己上传时那个文件名，不许看见服务器上的落盘位置。"""
    blob = json.dumps(response, ensure_ascii=False, default=str)
    assert str(tmp_path) not in blob, "响应里出现了本次上传的目录"
    for kind, item in _walk(response):
        if kind == "key":
            assert item not in _FORBIDDEN_KEYS, f"响应里出现了路径键 {item}"
        elif isinstance(item, str):
            assert _ABS_PATH.search(item) is None, f"响应里出现了绝对路径值：{item}"


def _assert_reading_not_a_claim(cell):
    """判据① 的下半句：这一格只报读数，不许把"没跑成"渲染成"这一页没有内容"。"""
    note = cell["degradation_note"]
    for claim in _EMPTINESS_CLAIMS:
        assert claim not in note, f"降级句把读数写成了断言：{claim}"


def _report(engine=None, path: Path = MIXED_PDF) -> loader.PdfExtractionReport:
    """独立对照：直接叫 loader 的账本，不走 HTTP、不走响应组装。"""
    with _quiet_parser_warnings():
        return loader.extract_pdf(str(path), engine=engine or PageAwareEngine(no_text_pages={5}))


def _upload_bytes(chat, filename: str, payload: bytes):
    from fastapi import UploadFile

    with _quiet_parser_warnings():
        return asyncio.run(
            chat.upload_document(
                UploadFile(filename=filename, file=io.BytesIO(payload)),
                classification=1,
                department="",
            )
        )


def _body(chars: int = 400) -> bytes:
    """一段够长、不会触发 R49 索引瘦身规则的正文，只用来让非 PDF 那几支走到正常返回。"""
    return ("费用报销管理制度正文。员工发生的差旅费须在三十日内提交审批。" * 40)[:chars].encode("utf-8")


# ------------------------------------------------------------------ 判据①⑥：端到端凭据


def test_the_upload_response_carries_the_page_source_book(monkeypatch, tmp_path):
    """判据①⑥：真走 ``upload_document``，「扫描页 3/6」这一串数字真的出现在响应 JSON 里。

    这里不替换 ``load_document``，也不直接调响应组装函数 —— 否则"扫描页 3/6"是桩话。
    """
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))

    response = _upload(chat, MIXED_PDF)

    cell = response[CELL_KEY]
    _assert_reading_or_nothing(cell)
    assert cell["page_count"] == MIXED_PAGE_COUNT
    assert cell["scanned_pages"] == len(MIXED_SCANNED_PAGES)
    assert cell["scanned_page_numbers"] == MIXED_SCANNED_PAGES
    assert cell["ocr_attempted"] is True
    assert cell["ocr_available"] is True
    assert cell["ocr_degraded_page_numbers"] == [], "本件这套桩里没有跑不成的页"
    # 逐页来源账：6 页四态齐现，且与 MIXED_PAGE_SHAPES 的逐页真值同一口径
    assert cell["source_counts"] == {
        PAGE_SOURCE_TEXT_LAYER: 2,
        PAGE_SOURCE_OCR: 2,
        PAGE_SOURCE_OCR_EMPTY: 1,
        PAGE_SOURCE_BLANK: 1,
    }
    # 数字真的过了 JSON 那一层（不是只有 Python 对象里对得上）
    assert set(cell) == _CELL_KEYS, "键集与契约那张表对不上"
    blob = json.dumps(response, ensure_ascii=False)
    assert CELL_KEY in blob
    assert json.loads(blob)[CELL_KEY] == cell
    assert json.loads(blob)[CELL_KEY]["scanned_pages"] == 3
    assert json.loads(blob)[CELL_KEY]["page_count"] == 6


def test_the_cell_moves_the_existing_book_without_a_second_vocabulary(monkeypatch, tmp_path):
    """判据①：那一格是"搬运"，词表与账本都还是 loader 现成的那一套。"""
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))
    report = _report()

    cell = _upload(chat, MIXED_PDF)[CELL_KEY]

    assert cell["source_counts"] == report.source_counts
    assert cell["scanned_page_numbers"] == list(report.scanned_page_numbers)
    assert cell["scanned_pages"] == report.scanned_pages
    assert set(cell["source_counts"]) <= _SOURCE_VOCABULARY
    assert cell["ocr_engine"] == report.ocr_engine
    assert cell["ocr_dpi"] == report.ocr_dpi


def test_pages_that_did_not_run_are_named_by_number_not_called_empty(monkeypatch, tmp_path):
    """判据① 的下半句：「哪几页没跑成」是页号，不是「这一页没有内容」。"""
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}, fail_pages={2}))

    cell = _upload(chat, MIXED_PDF)[CELL_KEY]

    assert cell["ocr_degraded_page_numbers"] == [2]
    assert cell["ocr_available"] is True, "引擎在，只是这一页没跑成 —— 别把两件事并成一格"
    assert cell["scanned_pages"] == 3, "有一页没跑成，扫描页读数不许缩水成 2"
    # 「没跑成」与「跑成了、图里确实没字」在账里仍是两个来源码
    assert cell["source_counts"][PAGE_SOURCE_OCR_DEGRADED] == 1
    assert cell["source_counts"][PAGE_SOURCE_OCR_EMPTY] == 1
    assert "第2页" in cell["degradation_note"]
    _assert_reading_not_a_claim(cell)


def test_a_pdf_without_scanned_pages_reports_real_zeros(monkeypatch, tmp_path):
    """判据③ 的另一半：是 PDF 就没扫描页 ⇒ 那一格有读数且读数是 0，与"不是 PDF"两回事。"""
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages=set()))

    response = _upload(chat, BLANK_PDF)
    cell = response[CELL_KEY]

    assert cell is not None, "PDF 就该有逐页读数，哪怕一页扫不出来"
    _assert_reading_or_nothing(cell)
    assert response["status"] == "skipped"
    assert set(cell) == _CELL_KEYS, "被排除的那条返回路径换了键集"
    assert cell["scanned_pages"] == 0
    assert cell["ocr_attempted"] is False, "没有扫描页时连 OCR 通道都不该开"
    assert cell["degradation_note"] == "", "没有降级就别造一句降级说明"
    assert cell["source_counts"] == {PAGE_SOURCE_BLANK: cell["page_count"]}


def test_the_reading_survives_a_real_http_round_trip(monkeypatch, tmp_path):
    """判据⑥ 最硬的一枚：真过 FastAPI 路由与 JSON 编码器，不是按协程直接调用。

    桩具与 ``tests/test_file_upload_security.py::_document_upload_client`` 同一套（不起端口、
    不起 app.main 的 lifespan）。这一枚还顺带过了 ``asyncio.to_thread`` 那一步：读数是
    工作线程产出的，要能回到事件循环里的响应组装，否则客户拿到的就是 null。
    """
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.agents import tools
    from app.api.v1 import chat

    _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))
    monkeypatch.setattr(tools, "rebuild_bm25", lambda: None)
    probe = FastAPI()
    probe.include_router(chat.router, prefix="/api/v1")
    client = TestClient(probe)

    with _quiet_parser_warnings():
        response = client.post(
            "/api/v1/upload",
            files={"file": (MIXED_PDF.name, MIXED_PDF.read_bytes(), "application/pdf")},
            data={"classification": "1", "department": ""},
        )

    assert response.status_code == 200, response.text
    body = response.text
    assert CELL_KEY in body and "3" in body
    cell = response.json()[CELL_KEY]
    _assert_reading_or_nothing(cell)
    assert (cell["scanned_pages"], cell["page_count"]) == (3, MIXED_PAGE_COUNT)
    assert cell["scanned_page_numbers"] == MIXED_SCANNED_PAGES
    assert cell["ocr_attempted"] is True and cell["ocr_available"] is True
    assert str(tmp_path) not in body, "过了一趟 HTTP 就把服务器路径带出去了"

# ------------------------------------------------------------------ 判据②：一把尺，两档话


def test_engine_unavailable_and_one_page_failing_are_two_different_answers(monkeypatch, tmp_path):
    """判据②：``ocr_available=False`` 与「引擎在、这一页没跑成」必须分得开（R298 定的口径）。"""
    # A 档：引擎不可用（R298 同一手法 —— 把模型文件名换成不存在的那一枚）
    chat_a = _wire(monkeypatch, tmp_path)
    monkeypatch.setattr(ocr_channel, "_MODEL_FILE_NAMES", ("models/absent.onnx",))
    missing = _upload(chat_a, MIXED_PDF)[CELL_KEY]

    assert missing["ocr_available"] is False
    assert missing["ocr_attempted"] is True, "通道开过又塌了，与压根没开是两件事"
    assert missing["degradation_note"].startswith(ocr_channel.ENGINE_UNAVAILABLE_NOTE)
    assert sorted(missing["ocr_degraded_page_numbers"]) == MIXED_SCANNED_PAGES

    # B 档：引擎在，只有第 2 页没跑成
    chat_b = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}, fail_pages={2}))
    partial = _upload(chat_b, MIXED_PDF)[CELL_KEY]

    assert partial["ocr_available"] is True
    assert partial["ocr_degraded_page_numbers"] == [2]
    assert partial["degradation_note"].startswith(loader.DEGRADATION_NOTE_PREFIX)
    assert missing["degradation_note"] != partial["degradation_note"]
    # 尺只有一把：两句都逐字等于 loader 自己算出来的那一句
    assert partial["degradation_note"] == _report(
        PageAwareEngine(no_text_pages={5}, fail_pages={2})
    ).degradation_sentence
    assert ocr_channel.ENGINE_UNAVAILABLE_NOTE != loader.DEGRADATION_NOTE_PREFIX



# ------------------------------------------------------------------ 判据③：非 PDF 那一格


@pytest.mark.parametrize("format_name", ["txt", "md", "docx"])
def test_a_non_pdf_upload_answers_a_whole_none_not_an_empty_object(monkeypatch, tmp_path, format_name):
    """判据③：``.txt`` / ``.md`` / ``.docx`` 的 ``pdf_extraction`` 整格 ``None``。

    ``.docx`` 也在这条里 —— 它有表格账，但**没有逐页来源账**，那一格说的是逐页，所以它
    照样是 None。判据点名禁止的形状是"回一个空对象冒充读过了"，这里逐字钉住。
    顺带钉住这一支的正文路径没被本单改坏：目录行 ready、正文真的交给了检索层。
    """
    retriever = RecordingRetriever()
    chat = _wire(monkeypatch, tmp_path, retriever)
    filename = f"policy.{format_name}"

    if format_name == "docx":
        source = _mini_docx(tmp_path / "in")
        response = _upload(chat, source)
    else:
        response = _upload_bytes(chat, filename, _body())

    assert CELL_KEY in response, "这一格在响应形状里缺席了"
    assert response[CELL_KEY] is None, "非 PDF 塞了东西：判据③ 点名要禁"
    _assert_reading_or_nothing(response[CELL_KEY])
    assert response["parse_status"] == "ready"
    assert response["status"] == "ok"
    assert retriever.add_calls == [filename if format_name != "docx" else "notes.docx"]


def test_the_two_loader_exits_agree_on_the_same_text(monkeypatch, tmp_path):
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))
    """判据③ 的后半句：新入口不许为这一格改坏任何一条正文通道 —— 同一个出口两个正文。"""
    with _quiet_parser_warnings():
        pdf = loader.extract_document_with_reports(str(MIXED_PDF))
        pdf_text = loader.load_document(str(MIXED_PDF))
    markdown = tmp_path / "in" / "note.md"
    markdown.parent.mkdir(parents=True, exist_ok=True)
    markdown.write_text("# 季度经营分析\n\n营收环比增长 12%。\n", encoding="utf-8")
    with _quiet_parser_warnings():
        md = loader.extract_document_with_reports(str(markdown))
        md_text = loader.load_document(str(markdown))
    docx = _mini_docx(tmp_path / "in2")
    with _quiet_parser_warnings():
        word = loader.extract_document_with_reports(str(docx))
        word_text = loader.load_document(str(docx))

    assert pdf.text == pdf_text and pdf.pdf is not None
    assert md.text == md_text and md.pdf is None
    assert word.text == word_text and word.pdf is None, ".docx 有表格账，但没有逐页账"
    assert md.tables.source == "", "非带账格式要写明「未参与」，不是「参与了一枚表没抽到」"
    assert md.tables.attached is False


def test_the_holdover_pass_produces_no_second_book():
    """新入口与 ``load_pdf`` 同一口径：重入闸里那一层交回被扣住的正文，不重复产账。"""
    key = loader._pdf_pass_key(str(MIXED_PDF))
    with _quiet_parser_warnings():
        with loader._pdf_prose_pass(key, "被扣住的那一份正文"):
            held = loader.extract_document_with_reports(str(MIXED_PDF))

    assert held.text == "被扣住的那一份正文"
    assert held.pdf is None, "重入那一层不该产第二本逐页账"


# ------------------------------------------------------------------ 判据④：读数不落库


def test_the_reading_is_not_written_into_the_catalog(monkeypatch, tmp_path):
    """判据④：这一格是"这一次上传的读数"，不是历史账 —— 一个字都不许落库。"""
    from app.documents import catalog

    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))
    recorded: list[dict] = []
    real_record = chat._record_uploaded_version

    def spy(**kwargs):
        result = real_record(**kwargs)
        recorded.append(dict(kwargs))
        return result

    monkeypatch.setattr(chat, "_record_uploaded_version", spy)

    response = _upload(chat, MIXED_PDF)

    assert response[CELL_KEY]["scanned_pages"] == 3, "先确认读数真的产生了，再钉它没落库"
    assert recorded, "一次目录行写入都没有，这半截钉子就是空钉"
    for kwargs in recorded:
        assert CELL_KEY not in kwargs
        assert not [key for key in kwargs if "extract" in key or "page" in key]
        for value in kwargs.values():
            assert not (isinstance(value, dict) and "page_count" in value), "读数进了落库参数"
    rows = catalog.list_document_versions(MIXED_PDF.name)
    assert rows, "目录行没写下"
    for row in rows:
        assert CELL_KEY not in row, "读数被写进了目录行"
        assert "page_count" not in row and "source_counts" not in row


def test_the_cell_brought_no_migration():
    """判据④ 的另一半：本单不落库 ⇒ ``migrations/`` 一个字都不该动（要留档是另一枚单）。"""
    offenders = []
    for script in sorted((REPO_ROOT / "migrations").glob("*.sql")):
        if CELL_KEY in script.read_text(encoding="utf-8", errors="ignore"):
            offenders.append(script.name)
    assert offenders == [], f"这一格被写进了迁移：{offenders}"


# ------------------------------------------------------------------ 不重解析（本单的代价口径）


def test_one_pdf_is_parsed_exactly_once(monkeypatch, tmp_path):
    """账来自产出它的那一次解析：为了这一格把 OCR 重跑一遍，本单不当这个价。"""
    chat = _wire(monkeypatch, tmp_path)
    engine = _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))
    parses: list[str] = []
    real_extract = loader.extract_pdf_with_tables

    def spy(file_path, **kwargs):
        parses.append(Path(str(file_path)).name)
        return real_extract(file_path, **kwargs)

    monkeypatch.setattr(loader, "extract_pdf_with_tables", spy)

    cell = _upload(chat, MIXED_PDF)[CELL_KEY]

    assert cell["scanned_pages"] == 3
    assert len(parses) == 1, f"一次上传解析了 {len(parses)} 遍 PDF"
    assert engine.calls == 3, "假引擎被叫的次数 = 扫描页数的两倍 ⇒ OCR 被重跑了"


def test_the_response_carries_no_absolute_path(monkeypatch, tmp_path):
    """判据⑤：整份响应（不只那一格）都不许出现服务器路径。"""
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))

    response = _upload(chat, MIXED_PDF)

    assert report_has_an_absolute_path(), "账本里本来就没有绝对路径？那这把钉子是空的"
    _assert_no_server_path(response, tmp_path)
    assert response["filename"] == MIXED_PDF.name, "客户能看到的只有他自己那个文件名"


def report_has_an_absolute_path() -> bool:
    """反证前提：``PdfExtractionReport.file_path`` 确实是一枚绝对路径 —— 泄露是真实风险。"""
    return bool(_ABS_PATH.search(_report().file_path))




# ------------------------------------------------------------------ 判据⑦：反证五把
# 每把都写清"摘了哪一把 ⇒ 红的正是哪一条"。摘完之后必须仍然看得见真读数，才算这有牙齿。


def test_counter_evidence_1_dropping_the_books_blanks_the_cell(monkeypatch, tmp_path):
    """①摘「新入口交回的账」：把出口退回不带账的那一枚 ⇒ 判据①⑥ 那些数字当场消失。

    这一把证明响应里的 3/6 是从账上来的，不是响应组装里写死的两个常量。
    """
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))
    with _quiet_parser_warnings():
        baseline = _upload(chat, MIXED_PDF)[CELL_KEY]
    assert baseline["scanned_pages"] == 3

    def no_books(file_path: str):
        with _quiet_parser_warnings():
            return loader._text_only_extraction(file_path, loader.load_document(file_path))

    monkeypatch.setattr(chat, "extract_document_with_reports", no_books)
    with _quiet_parser_warnings():
        emptied = _upload(chat, MIXED_PDF)[CELL_KEY]

    assert emptied is None, "账都摘了响应里还有读数 —— 那一格是硬编码的假读数"


def test_counter_evidence_2_a_fixed_degradation_sentence_is_caught(monkeypatch, tmp_path):
    """②摘「沿用那一把尺」：把 ``degradation_sentence`` 换成固定文案 ⇒ 判据② 那条等式红。

    同时证明这一格自己一个字都不造：尺改口，它也跟着改口。
    """
    engine_plan = lambda: PageAwareEngine(no_text_pages={5}, fail_pages={2})  # noqa: E731
    real_sentence = _report(engine_plan()).degradation_sentence
    assert real_sentence and "第2页" in real_sentence
    monkeypatch.setattr(
        loader.PdfExtractionReport,
        "degradation_sentence",
        property(lambda self: "文档已完整识别，没有任何降级"),
    )
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, engine_plan())

    cell = _upload(chat, MIXED_PDF)[CELL_KEY]

    assert cell["degradation_note"] == "文档已完整识别，没有任何降级"
    assert cell["degradation_note"] != real_sentence, "这一格私藏了另一套措辞（判据② 的病）"
    assert cell["ocr_degraded_page_numbers"] == [2], "固定文案抹不掉页号：读数与措辞是两回事"


def test_counter_evidence_3_an_empty_object_for_a_txt_is_rejected(monkeypatch, tmp_path):
    """③摘「非 PDF 整格 None」：让 ``.txt`` 也回一个空 dict ⇒ 判据③ 的形状尺判红。"""
    chat = _wire(monkeypatch, tmp_path)
    shipped = _upload_bytes(chat, "policy.txt", _body())[CELL_KEY]
    _assert_reading_or_nothing(shipped)
    assert shipped is None

    monkeypatch.setattr(chat, "_pdf_extraction_cell", lambda extraction: {})
    broken = _upload_bytes(chat, "policy.txt", _body())[CELL_KEY]

    with pytest.raises(AssertionError, match="空对象"):
        _assert_reading_or_nothing(broken)


def test_counter_evidence_4_a_leaked_absolute_path_is_caught(monkeypatch, tmp_path):
    """④摘「不透路径」：把 ``file_path`` 塞进那一格 ⇒ 判据⑤ 的扫描尺判红。

    账本里那枚路径是真的（``report.file_path`` 就是绝对路径），所以这一把不是稻草人。
    """
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))
    response = _upload(chat, MIXED_PDF)
    _assert_no_server_path(response, tmp_path)
    leaked = json.loads(json.dumps(response, ensure_ascii=False))
    leaked[CELL_KEY]["file_path"] = str(tmp_path / "stored" / "abcdef.pdf")

    with pytest.raises(AssertionError):
        _assert_no_server_path(leaked, tmp_path)

    smuggled = json.loads(json.dumps(response, ensure_ascii=False))
    smuggled[CELL_KEY]["degradation_note"] = str(tmp_path) + " 里有一页没跑成"
    with pytest.raises(AssertionError):
        _assert_no_server_path(smuggled, tmp_path)


def test_counter_evidence_5_a_claim_of_emptiness_is_caught(monkeypatch, tmp_path):
    """⑤摘「只报读数」：把降级句写成「这一页没有内容」⇒ 判据① 的引申尺判红。"""
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}, fail_pages={2}))
    cell = _upload(chat, MIXED_PDF)[CELL_KEY]
    _assert_reading_not_a_claim(cell)

    lying = dict(cell)
    lying["degradation_note"] = "扫描页 OCR 降级：第2页：这一页没有内容"
    with pytest.raises(AssertionError, match="断言"):
        _assert_reading_not_a_claim(lying)

    merged = dict(cell)
    # 三页扫描全记成空白页：总数仍然加得起来，红的是"自相矛盾"那一条，不是账不平
    merged["source_counts"] = {PAGE_SOURCE_BLANK: 6}
    with pytest.raises(AssertionError):
        _assert_reading_or_nothing(merged)


def test_counter_evidence_6_a_book_that_does_not_add_up_is_caught(monkeypatch, tmp_path):
    """⑥顺手一把：逐页来源账与总页数对不上 ⇒ 形状尺判红（这本账不能自己跟自己矛盾）。"""
    chat = _wire(monkeypatch, tmp_path)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))
    cell = _upload(chat, MIXED_PDF)[CELL_KEY]

    short = dict(cell)
    short["source_counts"] = {PAGE_SOURCE_TEXT_LAYER: 5}
    with pytest.raises(AssertionError):
        _assert_reading_or_nothing(short)

    zero = dict(cell)
    zero["page_count"] = 0
    with pytest.raises(AssertionError):
        _assert_reading_or_nothing(zero)


@pytest.mark.parametrize(
    "refusal,expected_reason",
    [
        ("文件内容未变化，已跳过", "unchanged_content"),
        ("向量库这一版不收", "index_refused"),
    ],
)
def test_the_index_refusal_paths_carry_the_same_reading(monkeypatch, tmp_path, refusal, expected_reason):
    """判据① 的形状承诺：索引层没收这份正文，读数照样交回来。

    R49 定过一条口径 —— 客户端不许靠"某个键不见了"去猜这次到底入没入索引。这一枚把那条例
    外扩到本单新增的这一格：四条返回路径给的是同一集键、同一份账。
    """
    from app.documents.index_policy import REASON_INDEX_REFUSED, REASON_UNCHANGED_CONTENT

    assert expected_reason in (REASON_INDEX_REFUSED, REASON_UNCHANGED_CONTENT)
    retriever = RecordingRetriever(result=(False, refusal))
    chat = _wire(monkeypatch, tmp_path, retriever)
    _engine_available(monkeypatch, PageAwareEngine(no_text_pages={5}))

    response = _upload(chat, MIXED_PDF)

    assert response["status"] == "skipped"
    cell = response[CELL_KEY]
    assert cell is not None, "索引层一拒绝，这一格就消失了"
    _assert_reading_or_nothing(cell)
    assert set(cell) == _CELL_KEYS, "拒收那条路径换了键集"
    assert cell["scanned_pages"] == 3 and cell["page_count"] == MIXED_PAGE_COUNT
    assert cell["source_counts"][PAGE_SOURCE_OCR] == 2
