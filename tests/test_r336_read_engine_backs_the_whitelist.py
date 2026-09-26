"""R336 · 收口名单里每一枚扩展名，脚下必须真有一枚可用的读引擎。

症状（总控在基点 `eec7ced` 实测，不是推测）：`app/api/v1/data.py` 的名单写着
`{".xlsx", ".xls", ".csv"}`，而 `app/tools/excel.py` 那条腿对 `.xls` 用的是
`engine="xlrd"` —— `xlrd` 既不在 `pyproject.toml`（只有 `openpyxl`），也不在这台机器的
`.venv` 里（现取 `importlib.util.find_spec("xlrd")` 回 None）。客户传一份真 `.xls` 上来，
走到的不是「这个格式我不支持，请另存为 .xlsx」，而是一条运行时 import 失败：更早一步还会
先在 openpyxl 的 `InvalidFileException` 上炸开，一路穿到上传路由的 `except Exception` 变 500。

本件的尺子不止治 `.xls` 一枚（判据 3）：它量「闸门放行的每一枚扩展名 ↔ 本机 import 得到的
引擎」这条不变量，而且每一格都用 `importlib.util.find_spec` 现读，不写死常量。下一位往名单里
加 `.numbers` / `.ods` 却忘了带引擎，它当场咬人 —— 两把反证（判据 4）就是这句话的凭据。

全部离线：不起服务、不连库、不打模型；夹具坐在 `tmp_path` + 文件形态的登记表上。
"""
from __future__ import annotations

import ast
import importlib.util
import re
from pathlib import Path
from typing import get_args

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PY = ROOT / "app" / "tools" / "excel.py"
DATA_PY = ROOT / "app" / "api" / "v1" / "data.py"

XLSX = "r336-sheet.xlsx"
CSV = "r336-table.csv"
XLS = "r336-legacy.xls"
TXT = "r336-notes.txt"
DEPT = "r336-finance"
OWNER = "r336-manager"

#: 一份真 .xls 的开头：Excel 97-2003 是 OLE2 复合文档，这个魔数就是 openpyxl 拒它的原因。
LEGACY_XLS_BYTES = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512

#: 把责任推回客户的黑话，一句都不许出现在拒绝里（判据 2）。
BLAME_JARGON = (
    "unsupported file type",
    "invalid file",
    "InvalidFileException",
    "ModuleNotFoundError",
    "ImportError",
    "internal_error",
    "系统异常",
    "未知错误",
    "内部错误",
    "请重试",
)


# ==================== 现读工具：三枚「白名单」对象与它们脚下的引擎 ====================


def _source(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _tree(path: Path) -> ast.Module:
    return ast.parse(_source(path))


def _module_assign(tree: ast.Module, name: str) -> ast.expr:
    """模块级 `NAME = ...` 的右值；找不到或有多处就抛（拿不准不许蒙）。"""
    values = [
        node.value
        for node in tree.body
        if isinstance(node, ast.Assign)
        for target in node.targets
        if isinstance(target, ast.Name) and target.id == name
    ]
    if len(values) != 1:
        raise AssertionError(f"{name} 应当恰好有一处模块级赋值，实取 {len(values)} 处")
    return values[0]


def _gate_extensions() -> set[str]:
    """闸门今天真正放行那一枚对象（`app/api/v1/data.py::DATA_FILE_EXTENSIONS`）。"""
    from app.api.v1 import data

    return {str(ext).lower() for ext in data.DATA_FILE_EXTENSIONS}


def _declared_extensions() -> set[str]:
    """产生上面那枚名单的声明表（`app/tools/excel.py::DATA_READ_ENGINES`）。"""
    from app.tools import excel

    return set(excel.accepted_data_file_extensions())


def _engineless_extensions() -> list[tuple[str, str]]:
    """判据 3 的尺子本体：名单里每一枚扩展名，它的引擎在本机 import 得到吗。

    三格设计上的讲究，缺一条这枚尺子就会撒谎：
      · 量的是**两枚对象**（闸门放行的 + 声明表里的）的并集 —— 反证把扩展名塞回声明表时，
        导入期快照还留在原地，只量快照的尺子会空响；
      · `find_spec` 由本件**自己现读**，不复用 `engine_is_importable()` 的结论 ——
        被测代码算错了可用性时，尺子必须还站得住；
      · 引擎写着 `None` 的那一格（pandas 内置通道）不在这里量，由 `test_..._round_trips`
        那枚真样本钉管住：声明「不需要额外的包」是要拿一份读得开的文件来还的。
    """
    from app.tools import excel

    offenders: list[tuple[str, str]] = []
    for ext in sorted(_gate_extensions() | _declared_extensions()):
        engine = excel.engine_for(ext)
        if engine is None:
            continue
        try:
            importable = importlib.util.find_spec(engine) is not None
        except (ImportError, ValueError, AttributeError):
            importable = False
        if not importable:
            offenders.append((ext, engine))
    return offenders


def _assert_the_admitted_set_has_engines() -> None:
    """反证窗要原样复用这一句，所以它是函数而不是内联 assert。"""
    offenders = _engineless_extensions()
    assert offenders == [], (
        "这些扩展名被收口名单放行，本机却没有它们的读引擎：%s —— 要么把依赖补上，"
        "要么把它搬去 REFUSED_DATA_FILE_READS 点名拒绝，不许留在名单里骗客户" % offenders
    )


def _assert_the_gate_matches_the_table() -> None:
    """两枚对象必须是同一份事实：闸门那张不许是导入期抄下来、之后各自漂的副本。"""
    drift = sorted(_gate_extensions() ^ _declared_extensions())
    assert drift == [], "闸门名单与引擎声明表已经分成两份：%s" % drift


def _assert_xls_is_refused_by_name() -> None:
    """判据 1 与判据 2 的交点：`.xls` 不在名单里，而且拒它的话是**点名说的**。

    只靠「这张表里没这一行」被挡住不算具名：那时候客户收到的是一句「不认识这个后缀」，
    而不是「这是 97-2003 老格式，另存为 .xlsx 再传」。
    """
    from app.tools import excel

    assert ".xls" not in excel.accepted_data_file_extensions(), ".xls 又回到收口名单了"
    assert ".xls" in excel.REFUSED_DATA_FILE_READS, (
        ".xls 没有具名拒绝那一行：它现在只靠「名单里不认识」被挡住，话说不到下一步做什么"
    )
    message = _refusal_message(XLS)
    assert "老格式" in message or "97-2003" in message, "这句拒绝没说清是哪一种老格式：%s" % message


def _refusal_message(name: str) -> str:
    from app.tools import excel

    refusal = excel.refuse_data_file(name)
    assert refusal is not None, "%s 居然被放行了：拒绝表那一格没生效" % name
    return str(refusal)


def _write_sample(directory: Path, filename: str) -> Path:
    """按后缀造一份**真能读**的文件：收口名单里每一枚扩展名都得有这样一份真样本。

    没有生成器就红 —— 这一格故意写得像门槛：「本机 import 得到引擎」是必要条件，
    不是充分条件。新格式要么带来一份读得开的样本，要么搬去具名拒绝那张表。
    """
    suffix = Path(filename).suffix.lower()
    target = directory / filename
    if suffix == ".csv":
        target.write_text("部门,月份,销售额\n财务,1月,100\n财务,2月,80\n", encoding="utf-8")
        return target
    if suffix == ".xlsx":
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(["部门", "月份", "销售额"])
        sheet.append(["财务", "1月", 100])
        sheet.append(["财务", "2月", 80])
        sheet.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
        sheet.cell(row=4, column=2, value="合并格测试")
        book.save(target)
        book.close()
        return target
    raise AssertionError(
        "收口名单里出现了 %s，但本件没有它的真样本生成器：要么补一枚，"
        "要么把这个格式搬去 REFUSED_DATA_FILE_READS" % suffix
    )


def _headers(username: str) -> dict[str, str]:
    from app.common.auth import create_token

    return {"Authorization": "Bearer %s" % create_token(username)}


def _user(username: str, department: str) -> dict[str, str]:
    return {"id": username, "username": username, "role": "manager", "department": department}


def _wire(monkeypatch, tmp_path):
    """DATA_DIR 与登记表搬到 tmp_path，账号表就地替换：一条真库连接都不发。"""
    from app.api.v1 import data
    from app.common import auth
    from app.storage import datasets
    from app.storage.datasets import DatasetRegistry

    registry = DatasetRegistry(root=tmp_path, metadata_path=tmp_path / "dataset-metadata.json")
    monkeypatch.setattr(data, "DATA_DIR", str(tmp_path))
    monkeypatch.setattr(data, "dataset_registry", registry)
    monkeypatch.setattr(datasets, "dataset_registry", registry)
    monkeypatch.setattr(auth, "get_user", lambda username: _user(username, DEPT))
    return data, registry


# ==================== 判据 3：那把通用尺子 ====================


def test_every_extension_the_gate_admits_has_an_importable_engine():
    _assert_the_admitted_set_has_engines()


def test_the_gate_and_the_engine_table_are_one_fact_not_two_lists():
    _assert_the_gate_matches_the_table()
    value = _module_assign(_tree(DATA_PY), "DATA_FILE_EXTENSIONS")
    assert isinstance(value, ast.Call), (
        "data.py 的收口名单又退回手抄字面量了：它必须是 accepted_data_file_extensions() 现算"
    )
    assert getattr(value.func, "id", "") == "accepted_data_file_extensions"


def test_the_xlsx_and_csv_legs_round_trip_a_real_file(tmp_path):
    """绿路径不许被治坏：名单里每一枚扩展名，真文件从磁盘走一遍要读得开。

    `.xlsx` 那一份特意带一个合并单元格，走的是 `_fill_merged_cells` 那条腿 —— 它正是旧实现
    比读腿更早撞上 openpyxl 的那一处，引擎拼错时它第一个塌。
    """
    from app.tools import excel

    for filename in sorted(excel.accepted_data_file_extensions()):
        frame = excel.load_excel(str(_write_sample(tmp_path, "walk" + filename)))
        assert len(frame.columns) == 3, filename
        assert len(frame) >= 2, filename


def test_a_merged_xlsx_still_fills_the_merged_span(tmp_path):
    from app.tools import excel

    frame = excel.load_excel(str(_write_sample(tmp_path, XLSX)))
    first = frame.columns[0]
    assert frame[first].tolist()[1] == frame[first].tolist()[2], "合并单元格的值没铺开"


def test_the_read_leg_never_hardcodes_an_engine_or_a_second_extension_branch():
    """旧病灶的语法形状是 `engine = "openpyxl" if ext == ".xlsx" else "xlrd"`：两处手抄。

    钉形状不钉结论：把 `xlrd` 写进任何字符串常量、或者在读腿里再抄一份后缀判断，都算复发。
    """
    tree = _tree(EXCEL_PY)
    xlrd_literals = [
        node.lineno
        for node in ast.walk(tree)
        if isinstance(node, ast.Constant) and node.value == "xlrd"
    ]
    assert xlrd_literals == [], "xlrd 又被当成一枚可用的引擎写进代码了：%s" % xlrd_literals

    load_excel = next(
        node
        for node in tree.body
        if isinstance(node, ast.FunctionDef) and node.name == "load_excel"
    )
    read_calls = [
        node
        for node in ast.walk(load_excel)
        if isinstance(node, ast.Call) and getattr(node.func, "attr", "") == "read_excel"
    ]
    assert len(read_calls) == 1, "读腿里冒出了第二处 read_excel：%s" % len(read_calls)
    engine_keyword = next((k for k in read_calls[0].keywords if k.arg == "engine"), None)
    assert engine_keyword is not None, "read_excel 不再声明引擎：它退回 pandas 的猜测就没人记账了"
    assert not isinstance(engine_keyword.value, ast.Constant), (
        "engine= 又写成了字面量：引擎名只能从 DATA_READ_ENGINES 现取"
    )
    literals = {
        node.value
        for node in ast.walk(load_excel)
        if isinstance(node, ast.Constant) and isinstance(node.value, str)
        and node.value.startswith(".")
    }
    assert literals == set(), "读腿里还有第二份后缀判断：%s" % sorted(literals)


# ==================== 判据 1 / 2：`.xls` 的具名拒绝 ====================


def test_xls_is_refused_by_name_and_not_folded_into_the_generic_answer():
    _assert_xls_is_refused_by_name()


def test_no_extension_is_both_admitted_and_refused():
    """两张表互斥：同一枚后缀既收又拒，等于两份口径互相打脸。"""
    from app.tools import excel

    overlap = sorted(set(excel.DATA_READ_ENGINES) & set(excel.REFUSED_DATA_FILE_READS))
    assert overlap == [], "这些后缀既在收口名单又在拒绝表里：%s" % overlap


def test_xls_hits_a_named_refusal_instead_of_a_runtime_import_failure(tmp_path):
    """旧症状的两副面孔都不许再出现：openpyxl 的 InvalidFileException、pandas 的 ImportError。"""
    from app.tools import excel

    sample = tmp_path / XLS
    sample.write_bytes(LEGACY_XLS_BYTES)

    with pytest.raises(excel.UnsupportedDataFile) as caught:
        excel.load_excel(str(sample))

    refusal = caught.value
    assert refusal.code == "unsupported_file"
    assert refusal.extension == ".xls"
    assert isinstance(refusal, ValueError), "形状变了：两条既有 except Exception 腿（agents/tools、alerts）会接不住"
    assert XLS in str(refusal)


def test_the_refusal_is_about_the_format_not_the_bytes(tmp_path):
    """后缀不合格时不必等到读磁盘才知道：一份不存在的 .xls 也该拿到同一句话。"""
    from app.tools import excel

    with pytest.raises(excel.UnsupportedDataFile) as caught:
        excel.ensure_data_file_readable(str(tmp_path / "nowhere" / XLS))
    assert "另存为" in str(caught.value)


@pytest.mark.parametrize("filename", [XLS, TXT, "r336-no-suffix"])
def test_every_refusal_sentence_names_a_next_step_an_employee_can_act_on(filename):
    """判据 2：拒绝必须可执行。「unsupported file type」这种把责任推回客户的话算不合格。"""
    message = _refusal_message(filename)

    assert "「%s」" % Path(filename).name in message, "这句拒绝没说清是哪一份文件：%s" % message
    assert "另存为" in message, "没有下一步动作的拒绝等于只说了坏消息：%s" % message
    assert re.search(r"另存为.{0,40}(xlsx|XLSX|csv|CSV)", message, re.S), (
        "「另存为」后面要点名一个真存在的目标格式，不然又是一句正确的废话：%s" % message
    )
    for jargon in BLAME_JARGON:
        assert jargon not in message, "句子里还留着推责黑话 %r：%s" % (jargon, message)


def test_the_generic_refusal_lists_the_formats_derived_not_handwritten():
    """不认识的后缀：把名单现算给客户看。名单变了这句话跟着变，才算同源。"""
    from app.tools import excel

    message = _refusal_message(TXT)
    for ext in sorted(excel.accepted_data_file_extensions()):
        assert ext in message, "名单里有 %s，这句话却没提它：%s" % (ext, message)


def test_the_refusal_code_is_ratified_and_r336_invents_no_new_code():
    """判据 1 的后半：错误码零新增。码 ∈ 封闭枚举，而且是从 `exc.code` 带上来的，不是现场手抄。"""
    from app.agents.contracts import ErrorEnvelope
    from app.tools import excel

    declared = set(get_args(ErrorEnvelope.model_fields["code"].annotation))
    assert excel.UnsupportedDataFile.code in declared
    assert excel.UnsupportedDataFile.code == "unsupported_file"

    tree = _tree(DATA_PY)
    refuse = next(
        node
        for node in tree.body
        if isinstance(node, ast.FunctionDef) and node.name == "_refuse_unreadable_data_file"
    )
    detail = None
    status_code = None
    for node in ast.walk(refuse):
        if isinstance(node, ast.Call) and getattr(node.func, "id", "") == "HTTPException":
            for keyword in node.keywords:
                if keyword.arg == "detail":
                    detail = keyword.value
                elif keyword.arg == "status_code":
                    status_code = keyword.value
    assert isinstance(detail, ast.Dict), "拒绝的响应形状不再是 {code, message}"
    keys = [k.value for k in detail.keys]
    assert keys == ["code", "message"], "响应键与既有信封形状脱钩：%s" % keys
    assert isinstance(detail.values[0], ast.Attribute), (
        "码名又在 data.py 里手抄了：它只能来自 UnsupportedDataFile.code 那一处"
    )
    assert isinstance(status_code, ast.Constant) and status_code.value == 400, (
        "拒绝不是 400：把「我不支持」报成 5xx 就是把产品口径伪装成故障"
    )


# ==================== 两条腿各自接线：上传与预览 ====================


def test_the_upload_route_refuses_xls_with_a_400_and_never_touches_the_disk(monkeypatch, tmp_path):
    """旧现场：`POST /upload-excel` 传真 .xls ⇒ 先落盘、再在 openpyxl 上炸、`except Exception`
    裸 re-raise ⇒ HTTP 500。客户看到的是「系统坏了」。今天：400 具名拒绝，且压根不落盘。
    """
    from fastapi.testclient import TestClient
    from app.main import app

    _data, registry = _wire(monkeypatch, tmp_path)

    response = TestClient(app, raise_server_exceptions=False).post(
        "/api/v1/upload-excel",
        headers=_headers(OWNER),
        files={"file": (XLS, LEGACY_XLS_BYTES, "application/vnd.ms-excel")},
    )

    assert response.status_code == 400, "拒绝又被报成故障（500）了：%s" % response.text
    detail = response.json()["detail"]
    assert isinstance(detail, dict) and sorted(detail) == ["code", "message"]
    assert detail["code"] == "unsupported_file"
    assert "另存为" in detail["message"] and ".xlsx" in detail["message"]
    assert not (tmp_path / XLS).exists(), "读不动的格式已经落进 DATA_DIR 了"
    assert registry.get_active_by_filename(XLS) is None
    assert list(tmp_path.glob(XLS)) == []


@pytest.mark.parametrize("filename", [XLSX, CSV])
def test_the_upload_route_still_reads_the_formats_it_admits(monkeypatch, tmp_path, filename):
    from fastapi.testclient import TestClient
    from app.main import app

    # 样本先躺在 DATA_DIR 之外的收件夹里：落到 DATA_DIR 的那一份必须由路由自己写。
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = _write_sample(inbox, filename)
    _data, registry = _wire(monkeypatch, tmp_path)

    response = TestClient(app, raise_server_exceptions=False).post(
        "/api/v1/upload-excel",
        headers=_headers(OWNER),
        files={"file": (filename, source.read_bytes(), "application/octet-stream")},
    )

    assert response.status_code == 200, response.text
    assert response.json()["filename"] == filename
    assert registry.get_active_by_filename(filename) is not None


def test_a_registered_legacy_xls_previews_as_a_named_refusal_not_a_500(monkeypatch, tmp_path):
    """登记表里可能留着一份切换前落盘的 .xls：预览那一格同样只能具名拒绝。

    放任它 ⇒ `except Exception` 兜成 500 `dataset_preview_failed`，把「这个格式我不读」
    说成「这台服务器出事了」—— 与本单要治的上传腿同一个病，另一张处方。
    """
    from fastapi.testclient import TestClient
    from app.agents.contracts import Principal
    from app.main import app

    _data, registry = _wire(monkeypatch, tmp_path)
    legacy = tmp_path / XLS
    legacy.write_bytes(LEGACY_XLS_BYTES)
    registry.register(legacy, principal=Principal.from_user(_user(OWNER, DEPT)), filename=XLS)

    response = TestClient(app, raise_server_exceptions=False).get(
        "/api/v1/data-files/%s/preview" % XLS, headers=_headers(OWNER)
    )

    assert response.status_code == 400, response.text
    assert response.json()["detail"]["code"] == "unsupported_file"
    assert "另存为" in response.json()["detail"]["message"]


def test_the_catalogue_stops_advertising_what_the_read_leg_cannot_open(monkeypatch, tmp_path):
    """本单明写承认的副作用：读不动的后缀不再出现在 /data-files 里。

    旧口径把 .xls 列出来、点开就 500，是「广告」了一枚没有腿的格式。现在列表与读腿同一张表：
    看不见它，也就不会被它骗一次。归属、密级、`restricted` 三格一概未动。
    """
    from fastapi.testclient import TestClient
    from app.agents.contracts import Principal
    from app.main import app

    _data, registry = _wire(monkeypatch, tmp_path)
    legacy = tmp_path / XLS
    legacy.write_bytes(LEGACY_XLS_BYTES)
    registry.register(legacy, principal=Principal.from_user(_user(OWNER, DEPT)), filename=XLS)

    listed = TestClient(app).get("/api/v1/data-files", headers=_headers(OWNER))
    assert listed.status_code == 200
    assert [row["filename"] for row in listed.json()["files"]] == []


# ==================== 判据 4：三把反证 ====================


def test_refutation_a_an_engineless_extension_pushed_back_into_the_whitelist_goes_red(monkeypatch):
    """临时把一枚没有引擎的扩展名塞回白名单 ⇒ 尺子必须红，并且点名是哪一枚。"""
    from app.tools import excel

    _assert_the_admitted_set_has_engines()
    _assert_the_gate_matches_the_table()

    monkeypatch.setitem(excel.DATA_READ_ENGINES, ".numbers", "r336_bogus_engine")

    assert ".numbers" in excel.accepted_data_file_extensions(), (
        "塞进声明表竟然没进白名单：白名单与声明表已经分成两份，尺子就量不到有人真加的那一格"
    )
    with pytest.raises(AssertionError) as caught:
        _assert_the_admitted_set_has_engines()
    assert ".numbers" in str(caught.value) and "r336_bogus_engine" in str(caught.value)

    # 尺子之外的两分本色：闸门快照立刻与声明表脱钩（第三枚牙），客户那一边仍然是一句话拒绝。
    with pytest.raises(AssertionError):
        _assert_the_gate_matches_the_table()
    refusal = excel.refuse_data_file("quarter.numbers")
    assert refusal is not None and "r336_bogus_engine" in str(refusal)
    assert ".numbers" in _refusal_message(TXT), "通用拒绝那句不再是现算名单"


def test_refutation_b_a_misspelled_xlsx_engine_goes_red(monkeypatch, tmp_path):
    """把 `.xlsx` 的引擎名拼错 ⇒ 必须红。这一把量的是「尺子读的是这台机器」而不是常量表。"""
    from app.tools import excel

    _assert_the_admitted_set_has_engines()

    monkeypatch.setitem(excel.DATA_READ_ENGINES, ".xlsx", "openpyxl_misspelled")

    with pytest.raises(AssertionError) as caught:
        _assert_the_admitted_set_has_engines()
    assert ".xlsx" in str(caught.value) and "openpyxl_misspelled" in str(caught.value)
    assert excel.unavailable_read_engines() == {".xlsx": "openpyxl_misspelled"}

    # 红的是真事：这份磁盘上躺着的 .xlsx 现在确实读不动了，而客户收到的仍是一句可执行的话。
    sample = _write_sample(tmp_path, XLSX)
    with pytest.raises(excel.UnsupportedDataFile) as refusal:
        excel.load_excel(str(sample))
    assert "openpyxl_misspelled" in str(refusal.value)
    assert "另存为" in str(refusal.value) and "管理员" in str(refusal.value)


def test_refutation_c_dropping_the_named_xls_row_goes_red(monkeypatch):
    """把 `.xls` 那行具名拒绝摘掉 ⇒ 红的必须是具名那一枚钉，别的格不许跟着溢。

    摘掉之后 `.xls` 仍然会被拒（走通用那句），也就是「看起来还是绿的」—— 正是这种
    「功能没塌所以没人发现」的格子，才需要一把专门指着措辞的牙。
    """
    from app.tools import excel

    _assert_xls_is_refused_by_name()

    monkeypatch.delitem(excel.REFUSED_DATA_FILE_READS, ".xls")

    with pytest.raises(AssertionError) as caught:
        _assert_xls_is_refused_by_name()
    assert "具名" in str(caught.value)

    # 不溢格：尺子、互斥、绿路径样本三枚照旧站得住。
    _assert_the_admitted_set_has_engines()
    _assert_the_gate_matches_the_table()
    assert sorted(set(excel.DATA_READ_ENGINES) & set(excel.REFUSED_DATA_FILE_READS)) == []
    assert ".xls" not in excel.accepted_data_file_extensions()
