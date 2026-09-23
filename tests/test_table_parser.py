"""PDF / DOCX structured table parsing.

Rows stay rows: the tests assert cell tuples survive exactly, that page and
table position are attached, and that open failures are reported as failed
rather than collapsed into an empty plain-text document.
"""
from app.rag.parser_result import BlockType, ParseStatus, ParserType
from app.rag.table_parser import extract_docx_tables, extract_pdf_tables


def _write_docx_with_tables(path):
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "month"
    table.cell(0, 1).text = "amount"
    table.cell(1, 0).text = "jan"
    table.cell(1, 1).text = "10"
    second = document.add_table(rows=1, cols=1)
    second.cell(0, 0).text = "footer table"
    document.save(str(path))


def _write_pdf_with_table(path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    with pdf.table() as table:
        row = table.row()
        row.cell("month")
        row.cell("amount")
        row = table.row()
        row.cell("jan")
        row.cell("10")
    pdf.output(str(path))


def test_docx_tables_preserve_rows_and_table_index(tmp_path):
    path = tmp_path / "report.docx"
    _write_docx_with_tables(path)

    run = extract_docx_tables(str(path))
    assert run.status is ParseStatus.SUCCESS
    assert len(run.blocks) == 2
    assert run.blocks[0].rows == (("month", "amount"), ("jan", "10"))
    assert run.blocks[0].block_type is BlockType.TABLE
    assert run.blocks[0].parser_type is ParserType.DOCX_TABLE
    assert run.blocks[0].location.table_index == 1
    assert run.blocks[1].location.table_index == 2


def test_docx_without_tables_is_empty(tmp_path):
    from docx import Document

    path = tmp_path / "prose.docx"
    document = Document()
    document.add_paragraph("just a paragraph")
    document.save(str(path))

    run = extract_docx_tables(str(path))
    assert run.status is ParseStatus.EMPTY
    assert run.blocks == ()


def test_docx_open_failure_is_failed(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip archive")

    run = extract_docx_tables(str(path))
    assert run.status is ParseStatus.FAILED
    assert run.errors[0].code == "docx_table_open_failed"


def test_pdf_tables_preserve_rows_and_page(tmp_path):
    path = tmp_path / "table.pdf"
    _write_pdf_with_table(path)

    run = extract_pdf_tables(str(path))
    assert run.status is ParseStatus.SUCCESS
    assert run.blocks[0].rows == (("month", "amount"), ("jan", "10"))
    assert run.blocks[0].parser_type is ParserType.PDF_TABLE
    assert run.blocks[0].location.page == 1
    assert run.blocks[0].location.table_index == 1


def test_pdf_without_tables_is_empty(tmp_path):
    from fpdf import FPDF

    path = tmp_path / "prose.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="plain paragraph")
    pdf.output(str(path))

    run = extract_pdf_tables(str(path))
    assert run.status is ParseStatus.EMPTY
    assert run.blocks == ()


def test_pdf_open_failure_is_failed(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf")

    run = extract_pdf_tables(str(path))
    assert run.status is ParseStatus.FAILED
    assert run.errors[0].code == "pdf_table_open_failed"


def test_pdf_partial_page_failure_is_partial_success(monkeypatch, tmp_path):
    class GoodPage:
        def extract_tables(self):
            return [[["month", "amount"], ["jan", "10"]]]

    class BrokenPage:
        def extract_tables(self):
            raise RuntimeError("synthetic page failure")

    class FakePdf:
        pages = [GoodPage(), BrokenPage()]

        def close(self):
            pass

    monkeypatch.setattr("pdfplumber.open", lambda path: FakePdf())

    run = extract_pdf_tables(str(tmp_path / "any.pdf"))
    assert run.status is ParseStatus.PARTIAL_SUCCESS
    assert len(run.blocks) == 1
    assert run.errors[0].code == "pdf_table_page_failed"
