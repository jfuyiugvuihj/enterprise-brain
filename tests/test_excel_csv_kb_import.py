"""Minimal knowledge-base import for Excel and CSV.

These tests lock the import shape: one block per worksheet / file, a header row,
data rows with an explicit row range, and source-file / sheet provenance.  They
also prove the cap is reported instead of silently dropping rows.
"""
from app.rag.excel_kb_parser import extract_csv_kb, extract_excel_kb
from app.rag.parser_result import BlockType, ParseStatus, ParserType
from app.rag.parser_result import render_block_for_index


def _write_workbook(path):
    from openpyxl import Workbook

    workbook = Workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales.append(["month", "amount"])
    sales.append(["jan", "10"])
    sales.append(["feb", "20"])
    sales.append([])
    summary = workbook.create_sheet("Summary")
    summary.append(["note"])
    summary.append(["total"])
    workbook.save(str(path))


def test_excel_import_preserves_sheets_headers_rows_and_ranges(tmp_path):
    path = tmp_path / "sales.xlsx"
    _write_workbook(path)

    run = extract_excel_kb(str(path))
    assert run.status is ParseStatus.SUCCESS
    assert len(run.blocks) == 2

    sales = run.blocks[0]
    assert sales.block_type is BlockType.DATA_ROWS
    assert sales.parser_type is ParserType.EXCEL_KB
    assert sales.header == ("month", "amount")
    assert sales.rows == (("jan", "10"), ("feb", "20"))
    assert sales.location.sheet == "Sales"
    assert sales.location.row_start == 2
    assert sales.location.row_end == 3
    assert sales.location.column_end == 2
    assert sales.location.file_path == str(path)

    summary = run.blocks[1]
    assert summary.location.sheet == "Summary"
    assert summary.header == ("note",)
    assert summary.rows == (("total",),)


def test_excel_import_reports_truncation_as_partial(tmp_path):
    path = tmp_path / "big.xlsx"
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Big"
    sheet.append(["n"])
    for index in range(15):
        sheet.append([index])
    workbook.save(str(path))

    run = extract_excel_kb(str(path), max_rows_per_sheet=10)
    assert run.status is ParseStatus.PARTIAL_SUCCESS
    assert len(run.blocks[0].rows) == 10
    assert run.metadata["sheet_rows"]["Big"] == 15
    assert any("10 of 15" in warning for warning in run.warnings)


def test_excel_open_failure_is_failed(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"PK\x03\x04not a workbook")

    run = extract_excel_kb(str(path))
    assert run.status is ParseStatus.FAILED
    assert run.errors[0].code == "excel_kb_open_failed"


def test_excel_index_text_keeps_sheet_and_structure(tmp_path):
    path = tmp_path / "sales.xlsx"
    _write_workbook(path)

    run = extract_excel_kb(str(path))
    text = render_block_for_index(run.blocks[0])
    assert "sheet=Sales" in text
    assert text.endswith("jan\t10\nfeb\t20")


def test_csv_import_preserves_header_rows_range_and_source(tmp_path):
    path = tmp_path / "sales.csv"
    path.write_text("month,amount\njan,10\nfeb,20\n", encoding="utf-8")

    run = extract_csv_kb(str(path))
    assert run.status is ParseStatus.SUCCESS
    block = run.blocks[0]
    assert block.parser_type is ParserType.CSV_KB
    assert block.header == ("month", "amount")
    assert block.rows == (("jan", "10"), ("feb", "20"))
    assert block.location.row_start == 2
    assert block.location.row_end == 3
    assert block.location.file_path == str(path)


def test_csv_import_reads_gbk_encoding(tmp_path):
    path = tmp_path / "legacy.csv"
    path.write_bytes("月份,金额\n一月,10\n".encode("gbk"))

    run = extract_csv_kb(str(path))
    assert run.status is ParseStatus.SUCCESS
    assert run.blocks[0].header == ("月份", "金额")
    assert run.blocks[0].rows == (("一月", "10"),)


def test_csv_explicit_wrong_encoding_is_failed(tmp_path):
    path = tmp_path / "legacy.csv"
    path.write_bytes("月份,金额\n".encode("gbk"))

    run = extract_csv_kb(str(path), encoding="utf-8")
    assert run.status is ParseStatus.FAILED
    assert run.errors[0].code == "csv_kb_open_failed"
