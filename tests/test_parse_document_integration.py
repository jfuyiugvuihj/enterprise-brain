"""Unified ``parse_document`` entry point and plain-text compatibility.

These tests prove the new entry works for every format the V2 line handles
while the old ``load_document`` / ``load_*`` behaviour stays intact for the
existing upload and preview callers.
"""
import pytest

from app.rag import loader
from app.rag.parser_result import BlockType, ParseStatus, ParserType


def _write_text_pdf(path, text="hello world"):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text=text)
    pdf.output(str(path))


def _write_image_pdf(path):
    from fpdf import FPDF
    from io import BytesIO
    from PIL import Image

    image = Image.new("RGB", (4, 4), (255, 255, 255))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)

    pdf = FPDF()
    pdf.add_page()
    pdf.image(buffer, x=0, y=0, w=10, h=10)
    pdf.output(str(path))


def test_parse_text_pdf_matches_legacy_plain_text(tmp_path):
    path = tmp_path / "doc.pdf"
    _write_text_pdf(path)

    result = loader.parse_document(str(path), include_tables=False)
    assert result.status is ParseStatus.SUCCESS
    assert result.blocks[0].block_type is BlockType.TEXT
    assert result.blocks[0].parser_type is ParserType.PYPDF
    assert result.blocks[0].location.page == 1
    assert result.plain_text == loader.load_pdf(str(path))


def test_parse_docx_with_table_keeps_text_and_table_separate(tmp_path):
    from docx import Document

    path = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("intro paragraph")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "month"
    table.cell(0, 1).text = "amount"
    table.cell(1, 0).text = "jan"
    table.cell(1, 1).text = "10"
    document.save(str(path))

    result = loader.parse_document(str(path))
    assert result.status is ParseStatus.SUCCESS
    assert result.plain_text == "intro paragraph"
    assert any(block.block_type is BlockType.TABLE for block in result.blocks)
    table_block = next(block for block in result.blocks if block.block_type is BlockType.TABLE)
    assert table_block.rows == (("month", "amount"), ("jan", "10"))
    assert "month\tamount" in result.to_index_text()


def test_parse_txt_and_md_stay_identical_to_legacy(tmp_path):
    txt = tmp_path / "notes.txt"
    txt.write_text("line one\nline two", encoding="utf-8")
    md = tmp_path / "notes.md"
    md.write_bytes("# 季度经营分析\n\n营收环比增长 12%。\n".encode("utf-8"))

    txt_result = loader.parse_document(str(txt))
    md_result = loader.parse_document(str(md))
    assert txt_result.status is ParseStatus.SUCCESS
    assert txt_result.plain_text == loader.load_txt(str(txt))
    assert md_result.status is ParseStatus.SUCCESS
    assert md_result.plain_text == loader.load_md(str(md))
    assert md_result.blocks[0].parser_type is ParserType.MARKDOWN


def test_parse_document_still_rejects_unknown_extension(tmp_path):
    path = tmp_path / "mystery.xyz"
    path.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported file format"):
        loader.parse_document(str(path))


def test_parse_xlsx_and_csv_as_knowledge_base(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["month", "amount"])
    sheet.append(["jan", "10"])
    xlsx = tmp_path / "sales.xlsx"
    workbook.save(str(xlsx))

    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("month,amount\njan,10\n", encoding="utf-8")

    xlsx_result = loader.parse_document(str(xlsx))
    csv_result = loader.parse_document(str(csv_path))
    assert xlsx_result.status is ParseStatus.SUCCESS
    assert xlsx_result.blocks[0].block_type is BlockType.DATA_ROWS
    assert xlsx_result.blocks[0].location.sheet == "Sales"
    assert xlsx_result.metadata["sheet_rows"]["Sales"] == 1
    assert csv_result.status is ParseStatus.SUCCESS
    assert csv_result.blocks[0].parser_type is ParserType.CSV_KB


def test_parse_pdf_ocr_without_engine_is_reported_not_silent(tmp_path):
    path = tmp_path / "doc.pdf"
    _write_text_pdf(path)

    result = loader.parse_document(str(path), include_ocr=True)
    assert result.status is ParseStatus.PARTIAL_SUCCESS
    codes = {error.code for error in result.errors}
    assert {"ocr_engine_missing", "ocr_renderer_missing"} <= codes


def test_parse_pdf_ocr_runs_scanned_pages_through_engine(tmp_path):
    from app.rag.ocr.engines import StubOcrEngine

    path = tmp_path / "scan.pdf"
    _write_image_pdf(path)

    result = loader.parse_document(
        str(path),
        include_ocr=True,
        ocr_engine=StubOcrEngine(default_text="recognised", confidence=0.9),
        render_page=lambda page: b"rendered",
    )
    ocr_blocks = [block for block in result.blocks if block.block_type is BlockType.OCR]
    assert len(ocr_blocks) == 1
    assert ocr_blocks[0].content == "recognised"
    assert ocr_blocks[0].location.page == 1
    assert "recognised" in result.to_index_text()
